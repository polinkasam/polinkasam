#!/usr/bin/env python3
"""Print table/column inventories for every schema source as Markdown.

Read-only. Sources are pali's own schema files; nothing is copied except
table and column names.
"""
import re
import sys
from pathlib import Path


def circe_xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    out = []
    for name in wb.sheetnames:
        ws = wb[name]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        cols = [str(c) for c in header if c is not None]
        out.append((name, cols))
    return out


def prisma(path):
    text = Path(path).read_text()
    models = re.findall(r"model\s+(\w+)\s*\{([^}]*)\}", text)
    out = []
    for name, body in models:
        cols = [
            l.split()[0]
            for l in body.strip().splitlines()
            if l.strip() and not l.strip().startswith(("@@", "//"))
        ]
        out.append((name, cols))
    return out


def sql(path):
    text = Path(path).read_text()
    tables = re.findall(
        r"create table\s+(?:if not exists\s+)?[\"`]?(?:\w+\.)?(\w+)[\"`]?\s*\(([^;]*?)\);",
        text,
        re.I | re.S,
    )
    out = []
    for name, body in tables:
        cols = [
            l.strip().split()[0].strip('",`')
            for l in body.splitlines()
            if l.strip()
            and not l.strip().lower().startswith(("primary", "foreign", "constraint", "unique", "check", "--", ")"))
        ]
        out.append((name, cols))
    return out


def md(title, source, tables):
    lines = [f"### {title}", "", f"Source: `{source}`", "", "| Table | Columns |", "| --- | --- |"]
    for name, cols in tables:
        lines.append(f"| {name} | {', '.join(cols)} |")
    return "\n".join(lines) + "\n"


SOURCES = {
    "Circe operations schema (Google Sheets / Excel)": (
        "/Users/polinkasam/Desktop/CLs/FIN/Circe.xlsx",
        circe_xlsx,
    ),
    "Finance Dashboard (Prisma on Postgres)": (
        "/Users/polinkasam/projects/apps/finance-dashboard/repos/lexware-finance-dashboard/prisma/schema.prisma",
        prisma,
    ),
    "Saint Suaire (Supabase)": (
        "/Users/polinkasam/projects/apps/arcana-echo/repos/arcana-echo/supabase-schema.sql",
        sql,
    ),
    "Leira (Supabase)": (
        "/Users/polinkasam/projects/apps/layra/repos/leira/supabase/schema.sql",
        sql,
    ),
    "Grant World / Circe graph export (Supabase SQL artifact)": (
        "/Users/polinkasam/projects/other/grant-world/artifacts/grant-world/supabase-schema.sql",
        sql,
    ),
}

if __name__ == "__main__":
    missing = 0
    for title, (path, fn) in SOURCES.items():
        if not Path(path).exists():
            print(f"### {title}\n\nSource missing: `{path}`\n")
            missing += 1
            continue
        print(md(title, path, fn(path)))
    sys.exit(1 if missing else 0)
